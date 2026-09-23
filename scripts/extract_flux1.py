#!/usr/bin/env python3
"""Extract the official DGFiP Flux 1 annex into a deterministic JSON mapping."""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "regulatory/official/dgfip/v3.2/20260430_Annexe-1-Flux-1-v1.2.xlsx"
TARGET = ROOT / "regulatory/flux1-mapping.v1.2.json"


def text(value):
    if value is None:
        return None
    return str(value).strip()


def split_rules(value):
    value = text(value)
    return [item.strip() for item in value.splitlines() if item.strip()] if value else []


def rows(sheet, syntax):
    extracted = []
    for values in sheet.iter_rows(min_row=6, values_only=True):
        identifier = text(values[0])
        if not identifier:
            continue
        labels = [text(values[index]) for index in range(2, 6)]
        if syntax == "UBL":
            roots = split_rules(values[6])
            path = text(values[7])
        else:
            roots = ["/CrossIndustryInvoice"]
            path = text(values[6])
        extracted.append(
            {
                "id": identifier,
                "kind": "group" if identifier.startswith("BG-") else "term",
                "cardinality": text(values[1]),
                "label": next((label for label in reversed(labels) if label), identifier),
                "hierarchy": [label for label in labels if label],
                "roots": roots,
                "path": path,
                "logicalType": text(values[8]),
                "length": text(values[9]),
                "codelist": text(values[10]),
                "managementRule": text(values[11]),
                "definition": text(values[12]),
                "usageNote": text(values[13]),
                "trajectory": text(values[14]),
                "genericRules": split_rules(values[15]),
                "syntaxRules": split_rules(values[16]),
                "profiles": {
                    "base": text(values[17]) == "X",
                    "full": text(values[18]) == "X",
                },
                "comment": text(values[19]),
            }
        )
    return extracted


def main():
    workbook = load_workbook(SOURCE, read_only=True, data_only=True)
    payload = {
        "$schema": "https://json-schema.org/draft/2020-12/schema",
        "title": "DGFiP Flux 1 mapping",
        "source": SOURCE.name,
        "sourceVersion": "1.2",
        "dgfipExternalSpecificationsVersion": "3.2",
        "publishedDate": "2026-04-30",
        "syntaxes": {
            "UBL": rows(workbook["FE - Flux 1 - UBL"], "UBL"),
            "CII": rows(workbook["FE - Flux 1 - CII"], "CII"),
        },
    }
    TARGET.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Wrote {TARGET}: UBL={len(payload['syntaxes']['UBL'])}, CII={len(payload['syntaxes']['CII'])}")


if __name__ == "__main__":
    main()
